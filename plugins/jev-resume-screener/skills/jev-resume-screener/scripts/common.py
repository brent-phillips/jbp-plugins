"""Shared helpers: file extraction, redaction, and text normalization.

Everything in this file runs locally. Nothing here calls a network service.
"""
import hashlib
import json
import re
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


# ---------- extraction ----------

def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if ext in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: {path.name}")


# ---------- redaction ----------

EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE = re.compile(r"(?:\+?\d{1,2}[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}")
URL = re.compile(r"(?:https?://|www\.)\S+|\b(?:linkedin|github)\.com/\S+", re.I)
STREET = re.compile(
    r"\b\d{1,6}\s+(?:[A-Z][a-z]+\s){1,3}(?:Street|St|Avenue|Ave|Road|Rd|Boulevard|Blvd|"
    r"Lane|Ln|Drive|Dr|Court|Ct|Way|Place|Pl|Circle|Cir|Parkway|Pkwy)\b\.?"
)
CITY_STATE_ZIP = re.compile(r"\b[A-Z][a-zA-Z .]+,\s*[A-Z]{2}\s+\d{5}(?:-\d{4})?\b")
NAME_LINE = re.compile(r"^[A-Z][A-Za-z'\-.]+(?:\s+[A-Z][A-Za-z'\-.]+){1,3}$")


def guess_name(text: str) -> str | None:
    """Most resumes start with the candidate's name on its own line."""
    for line in text.splitlines()[:5]:
        line = line.strip()
        if not line:
            continue
        candidate = re.split(r"\s*[|,•]\s*", line)[0].strip()
        if len(candidate) <= 40 and NAME_LINE.match(candidate):
            return candidate
        return None
    return None


def redact(text: str) -> tuple[str, str | None]:
    """Strip direct identifiers so Jev sees qualifications, not identity.

    Returns (redacted_text, guessed_name). The name stays local.
    """
    name = guess_name(text)
    out = text
    if name:
        out = out.replace(name, "[CANDIDATE]")
        for token in name.split():
            token = token.strip(".")
            if len(token) >= 3:
                out = re.sub(rf"\b{re.escape(token)}\b", "[CANDIDATE]", out)
    out = EMAIL.sub("[EMAIL]", out)
    out = URL.sub("[URL]", out)
    out = PHONE.sub("[PHONE]", out)
    out = STREET.sub("[ADDRESS]", out)
    out = CITY_STATE_ZIP.sub("[LOCATION]", out)
    return out, name


# ---------- normalization for comparisons ----------

PLACEHOLDER = re.compile(r"\[(?:CANDIDATE|EMAIL|URL|PHONE|ADDRESS|LOCATION)\]")


def normalize_words(text: str) -> list[str]:
    text = PLACEHOLDER.sub(" ", text.lower())
    return re.findall(r"[a-z0-9][a-z0-9+#.\-]*[a-z0-9+#]|[a-z0-9]", text)


def shingles(words: list[str], k: int) -> set[tuple[str, ...]]:
    if len(words) < k:
        return {tuple(words)} if words else set()
    return {tuple(words[i:i + k]) for i in range(len(words) - k + 1)}


def jaccard(a: set, b: set) -> float:
    if not a or not b:
        return 0.0
    return len(a & b) / len(a | b)


def sha(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


# ---------- io ----------

def load_json(path: Path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def save_json(path: Path, data) -> None:
    Path(path).write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
